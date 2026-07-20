from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Callable

from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


@dataclass(frozen=True, slots=True)
class DocumentChunk:
    text: str
    metadata: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def parse_pdf(file_path: str | Path) -> list[DocumentChunk]:
    path = Path(file_path)
    reader = PdfReader(path)
    return [
        DocumentChunk(
            text=text,
            metadata={"source": path.name, "file_type": ".pdf", "page": page_number},
        )
        for page_number, page in enumerate(reader.pages, start=1)
        if (text := (page.extract_text() or "").strip())
    ]


def parse_docx(file_path: str | Path) -> list[DocumentChunk]:
    path = Path(file_path)
    document = DocxDocument(path)
    chunks: list[DocumentChunk] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            chunks.append(
                DocumentChunk(
                    text=text,
                    metadata={
                        "source": path.name,
                        "file_type": ".docx",
                        "paragraph": paragraph_number,
                    },
                )
            )
    return chunks


def parse_xlsx(file_path: str | Path) -> list[DocumentChunk]:
    path = Path(file_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    chunks: list[DocumentChunk] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [
                "\t".join("" if value is None else str(value) for value in row).strip()
                for row in worksheet.iter_rows(values_only=True)
            ]
            text = "\n".join(row for row in rows if row)
            if text:
                chunks.append(
                    DocumentChunk(
                        text=text,
                        metadata={
                            "source": path.name,
                            "file_type": ".xlsx",
                            "sheet": worksheet.title,
                        },
                    )
                )
    finally:
        workbook.close()
    return chunks


def parse_pptx(file_path: str | Path) -> list[DocumentChunk]:
    path = Path(file_path)
    presentation = Presentation(path)
    chunks: list[DocumentChunk] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if texts:
            chunks.append(
                DocumentChunk(
                    text="\n".join(texts),
                    metadata={
                        "source": path.name,
                        "file_type": ".pptx",
                        "slide": slide_number,
                    },
                )
            )
    return chunks


PARSERS: dict[str, Callable[[str | Path], list[DocumentChunk]]] = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".pptx": parse_pptx,
}

text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)


def chunk_documents(source_chunks: list[DocumentChunk]) -> list[DocumentChunk]:
    chunks: list[DocumentChunk] = []
    for source_index, source_chunk in enumerate(source_chunks):
        for split_index, text in enumerate(text_splitter.split_text(source_chunk.text)):
            chunks.append(
                DocumentChunk(
                    text=text,
                    metadata={
                        **source_chunk.metadata,
                        "source_chunk": source_index,
                        "chunk_index": split_index,
                    },
                )
            )
    return chunks


def process_document(file_path: str | Path) -> list[DocumentChunk]:
    path = Path(file_path)
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        supported = ", ".join(sorted(PARSERS))
        raise ValueError(f"Unsupported file type. Supported types: {supported}")
    return chunk_documents(parser(path))

